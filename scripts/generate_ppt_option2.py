import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_option_2_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Palette Corporate (Deloitte, DXC, Enel)
    C_BG_DARK = RGBColor(15, 23, 42)        # #0F172A Slate 900
    C_BG_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC Slate 50
    C_WHITE = RGBColor(255, 255, 255)
    C_NAVY = RGBColor(30, 41, 59)          # #1E293B Slate 800
    C_CARD_BG = RGBColor(255, 255, 255)
    C_CARD_BORDER = RGBColor(226, 232, 240)# #E2E8F0
    C_BLUE = RGBColor(37, 99, 235)         # #2563EB Primary Blue (DXC style)
    C_BLUE_LIGHT = RGBColor(239, 246, 255) # #EFF6FF
    C_GREEN = RGBColor(16, 185, 129)       # #10B981 Emerald Green (Deloitte style)
    C_GREEN_LIGHT = RGBColor(236, 253, 245)# #ECFDF5
    C_ENEL_CYAN = RGBColor(6, 182, 212)    # #06B6D4 Smart City Cyan (Enel style)
    C_RED = RGBColor(239, 68, 68)          # #EF4444 Danger Red
    C_AMBER = RGBColor(245, 158, 11)       # #F59E0B Warning Amber
    C_TEXT_MAIN = RGBColor(15, 23, 42)     # #0F172A
    C_TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
    
    blank_layout = prs.slide_layouts[6]
    
    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="ALBO SICURO — PROPOSTA STRATEGICA CORPORATE"):
        # Category / Overline
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = C_BLUE
        
        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = C_TEXT_MAIN

    def add_speaker_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # ==========================================
    # SLIDE 1: COVER & URBAN VISION (ENEL & DXC)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, C_BG_DARK)
    
    # Corporate targets badge
    corp_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(4.8), Inches(0.45))
    corp_badge.fill.solid()
    corp_badge.fill.fore_color.rgb = RGBColor(30, 41, 59)
    corp_badge.line.color.rgb = C_ENEL_CYAN
    tf_cb = corp_badge.text_frame
    p_cb = tf_cb.paragraphs[0]
    p_cb.text = "FOCUS GIURIA: DELOITTE | DXC TECHNOLOGY | ENEL"
    p_cb.font.size = Pt(10)
    p_cb.font.bold = True
    p_cb.font.color.rgb = RGBColor(165, 243, 252)
    p_cb.alignment = PP_ALIGN.CENTER
    
    # Title
    t1_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.3))
    tf1 = t1_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "Albo Sicuro"
    p1.font.size = Pt(54)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    
    # Subtitle
    s1_box = s1.shapes.add_textbox(Inches(0.8), Inches(3.3), Inches(11), Inches(1.2))
    tfs1 = s1_box.text_frame
    tfs1.word_wrap = True
    ps1 = tfs1.paragraphs[0]
    ps1.text = "Privacy & Security Auditor per la Smart City e la Pubblica Amministrazione"
    ps1.font.size = Pt(22)
    ps1.font.color.rgb = RGBColor(203, 213, 225)
    
    # 3 Corporate Strategic Pillars
    pillars = [
        ("Deloitte Focus", "Trustworthy AI & Zero-Knowledge", C_GREEN),
        ("DXC Focus", "Enterprise Middleware API-first", C_BLUE),
        ("Enel Focus", "Smart City & Urban Data Protection", C_ENEL_CYAN)
    ]
    for idx, (p_tag, p_desc, p_col) in enumerate(pillars):
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 3.9), Inches(4.7), Inches(3.6), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(30, 41, 59)
        card.line.color.rgb = p_col
        card.line.width = Pt(1.5)
        
        tf_c = card.text_frame
        tf_c.margin_top = Inches(0.2)
        tf_c.margin_left = Inches(0.25)
        tf_c.margin_right = Inches(0.25)
        tf_c.word_wrap = True
        
        pt = tf_c.paragraphs[0]
        pt.text = p_tag.upper()
        pt.font.size = Pt(10)
        pt.font.bold = True
        pt.font.color.rgb = p_col
        pt.space_after = Pt(4)
        
        pd = tf_c.add_paragraph()
        pd.text = p_desc
        pd.font.size = Pt(12)
        pd.font.bold = True
        pd.font.color.rgb = C_WHITE
        
    # Footer info
    f1_box = s1.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.5), Inches(0.5))
    tf_f1 = f1_box.text_frame
    pf1 = tf_f1.paragraphs[0]
    pf1.text = "Campionato Universitario AI 2026 | DIGITA Academy — Napoli | Traccia 2: PRIVACY"
    pf1.font.size = Pt(11)
    pf1.font.color.rgb = RGBColor(148, 163, 184)
    
    add_speaker_notes(s1,
        "Buongiorno a tutti. La Smart City moderna si fonda sulla fiducia dei cittadini e sull'interoperabilità dei servizi digitali. "
        "Eppure, ogni giorno, oltre 7.900 Comuni italiani pubblicano online decine di atti amministrativi dove, per un banale errore umano, "
        "finiscono in rete diagnosi oncologiche, disabilità di minori, ISEE e IBAN. "
        "Abbiamo creato Albo Sicuro: l'auditor intelligente per la Smart City che previene, intercetta e bonifica i leak di dati personali "
        "nella PA prima che diventino danni irreversibili e pesanti sanzioni del Garante.")

    # ==========================================
    # SLIDE 2: IL RISCHIO ECONOMICO & REGOLATORIO (DELOITTE)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, C_BG_LIGHT)
    add_header(s2, "Il Costo dell'Errore Umano nella Trasparenza Pubblica", "GOVERNANCE, RISCHIO & COMPLIANCE (FOCUS DELOITTE)")
    
    stat_data = [
        ("€ 10.000 – € 50.000", "Sanzione Media per Singolo Atto", "Il Garante della Privacy non ammette l'errore materiale: la responsabilità del Titolare è oggettiva (violazione Art. 2-septies Codice Privacy e Art. 26 D.Lgs. 33/2013).", C_RED),
        ("15 Giorni vs Permanente", "Perdita Reale dell'Oblio", "L'affissione legale dura 15 giorni, ma i PDF restano scaricabili e indicizzati dai motori di ricerca per anni, creando esposizione permanente per l'ente.", C_AMBER),
        ("Zero Filtri Intelligenti", "Responsabilità sul Funzionario", "I funzionari comunali validano a mano centinaia di allegati complessi ogni giorno con software legacy, senza alcun layer di controllo automatico di conformità.", C_BLUE)
    ]
    for idx, (metric, subtitle, desc, bar_col) in enumerate(stat_data):
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 4.0), Inches(1.8), Inches(3.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)
        
        # Color top bar
        bar = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8 + idx * 4.0), Inches(1.8), Inches(3.7), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = bar_col
        bar.line.fill.background()
        
        tf_card = card.text_frame
        tf_card.margin_top = Inches(0.4)
        tf_card.margin_left = Inches(0.3)
        tf_card.margin_right = Inches(0.3)
        tf_card.word_wrap = True
        
        p_m = tf_card.paragraphs[0]
        p_m.text = metric
        p_m.font.size = Pt(21)
        p_m.font.bold = True
        p_m.font.color.rgb = bar_col
        p_m.space_after = Pt(8)
        
        p_st = tf_card.add_paragraph()
        p_st.text = subtitle
        p_st.font.size = Pt(15)
        p_st.font.bold = True
        p_st.font.color.rgb = C_TEXT_MAIN
        p_st.space_after = Pt(14)
        
        p_d = tf_card.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = C_TEXT_MUTED
        
    add_speaker_notes(s2,
        "La legge impone la trasparenza amministrativa, ma il GDPR vieta categoricamente la diffusione di dati sensibili. "
        "Oggi la responsabilità ricade interamente su singoli funzionari: basta una svista in una determina per buoni spesa "
        "per esporre l'indigenza di decine di famiglie. "
        "Per l'amministrazione questo non è solo un grave danno etico e reputazionale: significa contenziosi legali, "
        "danno erariale e sanzioni pesantissime da parte del Garante.")

    # ==========================================
    # SLIDE 3: LA SOLUZIONE DUAL-ENGINE (DXC & DELOITTE)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, C_BG_LIGHT)
    add_header(s3, "Albo Sicuro: Protezione Preventiva e Continua", "ARCHITETTURA A DOPPIO LIVELLO (GATEKEEPER & WATCHDOG)")
    
    modes = [
        ("1. Gatekeeper Preventivo (Pre-Rilascio)", "CONTROLLO QUALITÀ E BONIFICA PRIMA DELLA PUBBLICAZIONE", [
            ("Upload & Validazione Istantanea", "Il funzionario carica la bozza dell'atto; la piattaforma individua entità e frasi critiche."),
            ("Explainable AI & Precedenti", "L'algoritmo cita l'articolo di legge violato e associa il provvedimento sanzionatorio del Garante."),
            ("Vera Redazione Vettoriale", "Esportazione immediata del PDF corretto con 'omissis' irreversibili pronto per la firma digitale.")
        ], C_BLUE),
        ("2. Watchdog Continuo (Demone Asincrono)", "IL SOC DELLA PRIVACY H24 PER IL DPO E LA SMART CITY", [
            ("Monitoraggio Portali & Flussi Web", "Un demone indipendente scansiona in loop gli albi pretori comunali e i flussi di pubblicazione."),
            ("Incidental Management Istantaneo", "Rileva violazioni già online e notifica in tempo reale il DPO con indice di priorità."),
            ("Countdown Esposizione Rimanente", "Calcola i giorni residui di affissione all'albo per intervenire prima delle segnalazioni all'Autorità.")
        ], C_GREEN)
    ]
    for idx, (title, sub, bullets, accent) in enumerate(modes):
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 6.0), Inches(1.8), Inches(5.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1)
        
        strip = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1 + idx * 6.0), Inches(2.1), Inches(5.1), Inches(0.8))
        strip.fill.solid()
        strip.fill.fore_color.rgb = C_BLUE_LIGHT if idx == 0 else C_GREEN_LIGHT
        strip.line.color.rgb = accent
        tf_s = strip.text_frame
        tf_s.word_wrap = True
        ps = tf_s.paragraphs[0]
        ps.text = title
        ps.font.size = Pt(14)
        ps.font.bold = True
        ps.font.color.rgb = accent
        ps_sub = tf_s.add_paragraph()
        ps_sub.text = sub
        ps_sub.font.size = Pt(8.5)
        ps_sub.font.bold = True
        ps_sub.font.color.rgb = C_TEXT_MUTED
        
        b_box = s3.shapes.add_textbox(Inches(1.1 + idx * 6.0), Inches(3.1), Inches(5.1), Inches(3.3))
        tf_b = b_box.text_frame
        tf_b.word_wrap = True
        for b_idx, (b_t, b_d) in enumerate(bullets):
            pb = tf_b.paragraphs[0] if b_idx == 0 else tf_b.add_paragraph()
            pb.text = f"• {b_t}"
            pb.font.size = Pt(13)
            pb.font.bold = True
            pb.font.color.rgb = C_TEXT_MAIN
            pb.space_before = Pt(8)
            
            pbd = tf_b.add_paragraph()
            pbd.text = f"  {b_d}"
            pbd.font.size = Pt(11)
            pbd.font.color.rgb = C_TEXT_MUTED
            pbd.space_after = Pt(6)
            
    add_speaker_notes(s3,
        "Albo Sicuro interviene su due livelli: "
        "A monte, come Gatekeeper: prima della pubblicazione, il funzionario carica il documento; "
        "la piattaforma individua le criticità e genera con un clic la versione conforme pronta per la firma. "
        "A valle, come Watchdog Continuo: un demone indipendente sorveglia costantemente l'albo, "
        "segnala all'istante le falle sfuggite ai controlli manuali e calcola i giorni di esposizione per consentire una bonifica tempestiva.")

    # ==========================================
    # SLIDE 4: PRIVACY BY DESIGN & ZERO-KNOWLEDGE (DELOITTE & DXC)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, C_BG_LIGHT)
    add_header(s4, "Architettura 'Zero-Knowledge AI' & Trustworthy AI", "INGEGNERIA DELLA PRIVACY (FOCUS DELOITTE & DXC)")
    
    # Banner
    pill_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.6))
    pill_box.fill.solid()
    pill_box.fill.fore_color.rgb = C_NAVY
    pill_box.line.fill.background()
    tf_pb = pill_box.text_frame
    ppb = tf_pb.paragraphs[0]
    ppb.text = "ZERO DATA LEAKAGE: L'INTELLIGENZA ARTIFICIALE VEDE IL CONTESTO NORMATIVO, MAI L'IDENTITÀ REALE"
    ppb.font.size = Pt(10.5)
    ppb.font.bold = True
    ppb.font.color.rgb = C_WHITE
    ppb.alignment = PP_ALIGN.CENTER
    
    steps = [
        ("1. Ingestion Documentale", "PDF Amministrativo", "Estrazione del testo nativo con coordinate vettoriali esatte.", C_BLUE),
        ("2. Pseudonimizzazione", "Engine Locale Off-line", "Regex con omocodia + stdnum (CF, IBAN) + spaCy NER. I dati diventano [PERSONA_1], [CF_1].", C_BLUE),
        ("3. Reasoning & RAG", "Zero-Knowledge LLM", "L'LLM ragiona su token astratti e consulta la knowledge base delle ordinanze del Garante Privacy.", C_GREEN),
        ("4. True Redaction", "PyMuPDF Vettoriale", "Distruzione fisica e irreversibile dei dati dal codice binario del PDF. Non è un rettangolo grafico.", C_GREEN)
    ]
    for idx, (st_t, st_sub, st_d, st_col) in enumerate(steps):
        sbox = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 2.95), Inches(2.6), Inches(2.8), Inches(3.9))
        sbox.fill.solid()
        sbox.fill.fore_color.rgb = C_CARD_BG
        sbox.line.color.rgb = st_col
        sbox.line.width = Pt(1.5)
        
        sh = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8 + idx * 2.95), Inches(2.6), Inches(2.8), Inches(0.45))
        sh.fill.solid()
        sh.fill.fore_color.rgb = st_col
        sh.line.fill.background()
        tf_sh = sh.text_frame
        psh = tf_sh.paragraphs[0]
        psh.text = st_t
        psh.font.size = Pt(11)
        psh.font.bold = True
        psh.font.color.rgb = C_WHITE
        psh.alignment = PP_ALIGN.CENTER
        
        tf_sb = sbox.text_frame
        tf_sb.margin_top = Inches(0.6)
        tf_sb.margin_left = Inches(0.2)
        tf_sb.margin_right = Inches(0.2)
        tf_sb.word_wrap = True
        
        p_sub = tf_sb.paragraphs[0]
        p_sub.text = st_sub
        p_sub.font.size = Pt(13)
        p_sub.font.bold = True
        p_sub.font.color.rgb = C_TEXT_MAIN
        p_sub.space_after = Pt(10)
        
        p_desc = tf_sb.add_paragraph()
        p_desc.text = st_d
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = C_TEXT_MUTED
        
    add_speaker_notes(s4,
        "Il maggior rischio nell'applicare l'IA alla privacy è inviare dati di cittadini a provider cloud esterni. "
        "Noi abbiamo applicato un'architettura rigorosa di Privacy by Design e Trustworthy AI: prima dell'analisi, "
        "un motore locale sostituisce ogni dato identificativo con token astratti. "
        "L'IA vede solo che [PERSONA_1] beneficia di un sussidio per grave disabilità e deduce la non conformità senza mai conoscere l'identità del cittadino. "
        "E per la bonifica finale, usiamo redazione vettoriale irreversibile: il dato sensibile cessa fisicamente di esistere all'interno del file.")

    # ==========================================
    # SLIDE 5: LIVE DEMO & AUDIT TRAIL (DELOITTE)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, C_BG_LIGHT)
    add_header(s5, "Dall'Allarme alla Risoluzione in 3 Clic", "ACCOUNTABILITY & AUDIT TRAIL (FOCUS DELOITTE)")
    
    demo_cards = [
        ("Fase 1: Alert Decisionale", "LIVE MONITORING FEED", "Il demone o l'operatore intercettano una determina per buoni spesa con ISEE ed esiti in chiaro. Semaforo Rosso immediato.", C_RED, "Status: Non Conforme (Gravità Alta)"),
        ("Fase 2: RAG sui Precedenti", "EXPLAINABLE COMPLIANCE", "L'IA cita l'Art. 26 c. 4 D.Lgs. 33/2013 e abbina il caso reale del Comune di Tricase (sanzione comminata dal Garante: € 15.000).", C_AMBER, "Precedente Storico: GPDP 9712044"),
        ("Fase 3: Bonifica & Sostituzione", "REDACTION SIDE-BY-SIDE", "Generazione istantanea del PDF corretto con omissis. Download immediato e tracciamento dell'evento di risoluzione nei log di audit.", C_GREEN, "Azioni: Risolvi / Sostituisci Atto")
    ]
    for idx, (dc_t, dc_sub, dc_d, dc_col, badge) in enumerate(demo_cards):
        dcard = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 4.0), Inches(1.8), Inches(3.7), Inches(4.8))
        dcard.fill.solid()
        dcard.fill.fore_color.rgb = C_CARD_BG
        dcard.line.color.rgb = C_CARD_BORDER
        dcard.line.width = Pt(1)
        
        bdg = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0 + idx * 4.0), Inches(2.0), Inches(3.3), Inches(0.45))
        bdg.fill.solid()
        bdg.fill.fore_color.rgb = C_BG_LIGHT
        bdg.line.color.rgb = dc_col
        tf_bd = bdg.text_frame
        p_bd = tf_bd.paragraphs[0]
        p_bd.text = badge
        p_bd.font.size = Pt(9.5)
        p_bd.font.bold = True
        p_bd.font.color.rgb = dc_col
        p_bd.alignment = PP_ALIGN.CENTER
        
        tf_dc = dcard.text_frame
        tf_dc.margin_top = Inches(0.9)
        tf_dc.margin_left = Inches(0.3)
        tf_dc.margin_right = Inches(0.3)
        tf_dc.word_wrap = True
        
        p_dct = tf_dc.paragraphs[0]
        p_dct.text = dc_t
        p_dct.font.size = Pt(16)
        p_dct.font.bold = True
        p_dct.font.color.rgb = C_TEXT_MAIN
        p_dct.space_after = Pt(6)
        
        p_dcsub = tf_dc.add_paragraph()
        p_dcsub.text = dc_sub
        p_dcsub.font.size = Pt(10)
        p_dcsub.font.bold = True
        p_dcsub.font.color.rgb = C_TEXT_MUTED
        p_dcsub.space_after = Pt(14)
        
        p_dcd = tf_dc.add_paragraph()
        p_dcd.text = dc_d
        p_dcd.font.size = Pt(12)
        p_dcd.font.color.rgb = C_TEXT_MUTED
        
    add_speaker_notes(s5,
        "Ecco Albo Sicuro in azione: il nostro demone intercetta una determina con i beneficiari dei buoni spesa e relativi ISEE. "
        "Il sistema accende il semaforo rosso e richiama dalla knowledge base il precedente ufficiale del Garante: "
        "'Attenzione, per un caso analogo il Comune di Tricase è stato sanzionato per 15.000 euro'. "
        "Il responsabile visualizza i passaggi critici evidenziati e, con un singolo clic, scarica la versione bonificata pronta per la sostituzione, "
        "sanando l'incidente in pochi secondi e registrando l'evento nell'audit trail.")

    # ==========================================
    # SLIDE 6: B2G MODEL & SYSTEM INTEGRATION (DXC & DELOITTE)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, C_BG_LIGHT)
    add_header(s6, "Modello B2G e Opportunità di Integrazione Enterprise", "SYSTEM INTEGRATION & GTM (FOCUS DXC & DELOITTE)")
    
    pillars_b2g = [
        ("Middleware API Plug-and-Play", "Add-on per i Grandi System Integrator", "Non rimpiazziamo i gestionali documentali esistenti (Maggioli, Halley, Engineering): ci integriamo come motore di certificazione e sicurezza per i contratti quadro PA gestiti da partner come DXC.", C_BLUE),
        ("Privacy-as-a-Service (PaaS)", "Supporto Strutturato al DPO", "Fornisce ai DPO e agli auditor di Deloitte una console centralizzata con report crittografici SHA-256, retention automatica e piena conformità al principio di accountability (Art. 5.2 GDPR).", C_GREEN),
        ("ROI Matematico per l'Ente", "Riduzione del Rischio Sanzionatorio", "Il canone annuale del servizio costa meno di una frazione di una singola sanzione del Garante (€ 10k-50k) o dei costi legali e risarcitori di un contenzioso cittadino.", C_AMBER)
    ]
    for idx, (p_t, p_sub, p_d, p_c) in enumerate(pillars_b2g):
        pcard = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 4.0), Inches(1.8), Inches(3.7), Inches(4.8))
        pcard.fill.solid()
        pcard.fill.fore_color.rgb = C_CARD_BG
        pcard.line.color.rgb = C_CARD_BORDER
        pcard.line.width = Pt(1)
        
        bar = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8 + idx * 4.0), Inches(1.8), Inches(3.7), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = p_c
        bar.line.fill.background()
        
        tf_pc = pcard.text_frame
        tf_pc.margin_top = Inches(0.4)
        tf_pc.margin_left = Inches(0.3)
        tf_pc.margin_right = Inches(0.3)
        tf_pc.word_wrap = True
        
        p1 = tf_pc.paragraphs[0]
        p1.text = p_t
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = p_c
        p1.space_after = Pt(8)
        
        p2 = tf_pc.add_paragraph()
        p2.text = p_sub
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_MAIN
        p2.space_after = Pt(14)
        
        p3 = tf_pc.add_paragraph()
        p3.text = p_d
        p3.font.size = Pt(11.5)
        p3.font.color.rgb = C_TEXT_MUTED
        
    add_speaker_notes(s6,
        "Il nostro modello non richiede ai Comuni di cambiare software gestionale: Albo Sicuro è progettato come middleware API integrabile "
        "nelle soluzioni che grandi player come DXC già forniscono alla PA. "
        "Per una qualsiasi amministrazione il ritorno sull'investimento è immediato: evitare anche una sola sanzione da 20.000 euro o un contenzioso legale "
        "ripaga il canone del servizio per diversi anni, garantendo ai vertici amministrativi la serenità dell'accountability GDPR.")

    # ==========================================
    # SLIDE 7: ROADMAP SMART CITY & CLOSING (ENEL)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, C_BG_DARK)
    
    t7_box = s7.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.0))
    tf7 = t7_box.text_frame
    p7 = tf7.paragraphs[0]
    p7.text = "Dall'Albo Pretorio alla Sicurezza dei Dati nella Smart City"
    p7.font.size = Pt(30)
    p7.font.bold = True
    p7.font.color.rgb = C_WHITE
    
    r_steps = [
        ("Q3 2026", "OCR Locale On-Premise", "Estensione dell'ingestion ad atti storici cartacei e scansioni senza inviare immagini a endpoint esterni.", C_BLUE),
        ("Q4 2026 (Focus Enel)", "Auditing IoT & Smart Metering", "Estensione della pipeline core ai flussi di telemetria urbana, log di smart meter e registri di videosorveglianza urbana.", C_ENEL_CYAN),
        ("Compliance Nazionale", "Qualificazione ACN Cloud PA", "Ospitalità su infrastruttura sovereign cloud certificata dall'Agenzia per la Cybersicurezza Nazionale.", C_GREEN)
    ]
    for idx, (time_b, r_title, r_desc, r_col) in enumerate(r_steps):
        rcard = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 4.0), Inches(2.2), Inches(3.7), Inches(3.2))
        rcard.fill.solid()
        rcard.fill.fore_color.rgb = RGBColor(30, 41, 59)
        rcard.line.color.rgb = r_col
        rcard.line.width = Pt(1.5)
        
        tf_rc = rcard.text_frame
        tf_rc.margin_top = Inches(0.3)
        tf_rc.margin_left = Inches(0.25)
        tf_rc.margin_right = Inches(0.25)
        tf_rc.word_wrap = True
        
        pr1 = tf_rc.paragraphs[0]
        pr1.text = time_b
        pr1.font.size = Pt(13)
        pr1.font.bold = True
        pr1.font.color.rgb = r_col
        pr1.space_after = Pt(6)
        
        pr2 = tf_rc.add_paragraph()
        pr2.text = r_title
        pr2.font.size = Pt(13)
        pr2.font.bold = True
        pr2.font.color.rgb = C_WHITE
        pr2.space_after = Pt(10)
        
        pr3 = tf_rc.add_paragraph()
        pr3.text = r_desc
        pr3.font.size = Pt(11)
        pr3.font.color.rgb = RGBColor(203, 213, 225)
        
    tcard = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.1))
    tcard.fill.solid()
    tcard.fill.fore_color.rgb = RGBColor(30, 41, 59)
    tcard.line.color.rgb = C_ENEL_CYAN
    tf_tc = tcard.text_frame
    tf_tc.margin_left = Inches(0.3)
    tf_tc.margin_top = Inches(0.2)
    ptc1 = tf_tc.paragraphs[0]
    ptc1.text = "Team di Sviluppo Albo Sicuro — Campionato Universitario AI (DIGITA Academy / AI2B)"
    ptc1.font.size = Pt(14)
    ptc1.font.bold = True
    ptc1.font.color.rgb = C_ENEL_CYAN
    ptc2 = tf_tc.add_paragraph()
    ptc2.text = "La vera digitalizzazione di una Smart City tutela i dati dei cittadini mentre rende i servizi più trasparenti e intelligenti. Grazie!"
    ptc2.font.size = Pt(11)
    ptc2.font.color.rgb = RGBColor(203, 213, 225)
    
    add_speaker_notes(s7,
        "La vera digitalizzazione di una Smart City non consiste nel pubblicare indiscriminatamente tutto online, "
        "ma nel tutelare i dati dei cittadini mentre si rende la città più trasparente e intelligente. "
        "Albo Sicuro trasforma la compliance da rischio paralizzante a garanzia automatica e scalabile per la Smart City. "
        "Grazie per l'attenzione, siamo aperti alle vostre domande!")

    # Save
    out_path = os.path.join(os.getcwd(), "Albo_Sicuro_Presentazione_Opzione2.pptx")
    prs.save(out_path)
    print(f"Presentation saved successfully to: {out_path}")

if __name__ == '__main__':
    create_option_2_presentation()
